import pandas as pd
from datetime import date

import openpyxl
import pytest

from sigma_conn import negocio as N
from sigma_conn import objetivos as O
from sigma_conn import objetivos_excel as OX


# ----------------------------------------------------------------------------- calendario
def test_dias_de_venta_lunes_a_sabado():
    assert O.dias_de_venta_mes(2026, 9) == 26      # septiembre 2026: 30 días, 4 domingos
    assert O.dias_de_venta_mes(2026, 10) == 27     # octubre 2026: 31 días, 4 domingos
    assert O.dias_de_venta(date(2026, 9, 6), date(2026, 9, 6)) == 0   # domingo
    assert O.dias_de_venta(date(2026, 9, 7), date(2026, 9, 5)) == 0   # rango invertido


def test_transcurridos_mas_restantes_es_el_mes():
    f = date(2026, 9, 28)
    assert O.dias_transcurridos(f) == 24 and O.dias_restantes(f) == 2
    assert O.dias_transcurridos(f) + O.dias_restantes(f) == O.dias_de_venta_mes(2026, 9)
    assert O.dias_restantes(date(2026, 9, 30)) == 0


def test_objetivo_diario_es_mensual_sobre_26():
    assert O.objetivo_diario(52_000_000) == pytest.approx(2_000_000)


# ----------------------------------------------------------------------------- facturación
def test_evaluar_facturacion_sin_escalon_pero_proyecta_el_primero():
    r = O.evaluar_facturacion(48_000_000, N.PERFIL_GENERAL, date(2026, 9, 28))
    assert r["escalon"] == 0 and r["premio"] == 0
    assert r["ritmo_diario"] == pytest.approx(2_000_000)
    assert r["proyeccion"] == pytest.approx(52_000_000)           # 2 M$ por día * 26 días de venta
    e1 = r["escalones"][0]
    assert e1["faltante"] == pytest.approx(2_000_000) and e1["media_necesaria"] == pytest.approx(1_000_000)
    assert e1["proyeccion_alcanza"] and not r["escalones"][1]["proyeccion_alcanza"]
    assert r["premio_proyectado"] == 200_000 and r["siguiente"]["escalon"] == 1


def test_evaluar_facturacion_escalon_y_premio_en_pesos():
    r = O.evaluar_facturacion(61_000_000, N.PERFIL_GENERAL, date(2026, 9, 28))
    assert r["escalon"] == 2 and r["premio"] == 400_000
    assert r["siguiente"]["escalon"] == 3 and r["siguiente"]["faltante"] == pytest.approx(11_000_000)
    assert O.evaluar_facturacion(72_000_000, N.PERFIL_GENERAL, date(2026, 9, 28))["premio"] == 600_000
    assert O.evaluar_facturacion(72_000_000, N.PERFIL_GENERAL, date(2026, 9, 28))["siguiente"] is None


def test_evaluar_facturacion_usa_la_escala_del_perfil():
    aass = O.evaluar_facturacion(110_000_000, N.PERFIL_AASS, date(2026, 9, 28))     # 100/120/144
    assert aass["escalon"] == 1 and aass["premio"] == 200_000
    interior = O.evaluar_facturacion(70_000_000, N.PERFIL_INTERIOR, date(2026, 9, 28))  # 60/70/85
    assert interior["escalon"] == 2 and interior["premio"] == 400_000


def test_evaluar_facturacion_ultimo_dia_sin_dias_restantes():
    r = O.evaluar_facturacion(40_000_000, N.PERFIL_GENERAL, date(2026, 9, 30))
    assert r["dias_restantes"] == 0 and r["escalones"][0]["media_necesaria"] is None   # ya no hay forma de llegar
    ok = O.evaluar_facturacion(50_000_000, N.PERFIL_GENERAL, date(2026, 9, 30))
    assert ok["escalones"][0]["media_necesaria"] == 0.0


def test_avance():
    assert O.avance(50, 200) == 0.25 and O.avance(10, 0) is None


# ----------------------------------------------------------------------------- configuración de negocio
def test_escalas_crecientes_y_premios():
    for escalas in N.ESCALAS_FACTURACION.values():
        assert list(escalas) == sorted(escalas) and len(escalas) == len(N.PREMIOS_ESCALON) == 3
    assert N.ESCALAS_FACTURACION["GENERAL"][0] == 50_000_000   # no existe el escalón de 25 M$


def test_mix_por_marca_suma_100():
    for mix in N.MIX_MARCA.values():
        assert sum(mix.values()) == pytest.approx(1.0) and set(mix) == set(N.MARCAS)


def test_marca_de_proveedor():
    assert N.marca_de_proveedor("UNILEVER DE ARGENTINA SA") == "UNILEVER"
    assert N.marca_de_proveedor("GRUPO PEÑAFLOR S.A") == "PEÑAFLOR"
    assert N.marca_de_proveedor("GRUPO YUANFEN SRL (PALAU)") == "PALAU"
    assert N.marca_de_proveedor("DEPOSITO SALTA") is None and N.marca_de_proveedor(None) is None


def test_vendedores_con_perfil_tienen_supervisor_y_fili_va_con_mauro():
    assert set(N.VENDEDOR_PERFIL) <= set(N.SUPERVISOR_VENDEDOR)
    assert N.SUPERVISORES[N.SUPERVISOR_VENDEDOR["122"]] == "AMAYA MAURO"
    assert N.VENDEDOR_PERFIL["100"] == N.VENDEDOR_PERFIL["119"] == "AASS"
    assert N.VENDEDOR_PERFIL["120"] == N.VENDEDOR_PERFIL["111"] == "INTERIOR"


def test_grupos_de_cobertura_no_se_pisan():
    todos = [v for g in N.COBERTURA_GRUPOS.values() for v in g]
    assert len(todos) == len(set(todos)) == 13 and len(N.COBERTURA_GRUPOS["PREVENTA"]) == 9


def test_tablas_de_configuracion():
    f = O.build_obj_facturacion()
    assert len(f) == 9 and f.loc[(f.perfil == "AASS") & (f.escalon == 3), "premio"].item() == 600_000
    assert len(O.build_obj_mix_marca()) == 12
    vp = O.build_cfg_vendedor_perfil()
    assert set(vp["vendedor_id"]) == set(N.VENDEDOR_PERFIL) and vp.loc[vp.vendedor_id == "122", "supervisor"].item() == "AMAYA MAURO"
    assert len(O.build_cfg_supervisor_vendedor()) == len(N.SUPERVISOR_VENDEDOR)


# ----------------------------------------------------------------------------- lectura del Excel (libro sintético)
def _libro_sintetico():
    wb = openpyxl.Workbook()
    wb.active.title = "OBJ_COBERTURA"
    ws = wb["OBJ_COBERTURA"]
    for r, (cat, total) in zip((7, 8, 9), (("BPC", 1398), ("FOOD ", 1719), ("HC", 1680))):
        ws[f"B{r}"], ws[f"C{r}"], ws[f"G{r}"], ws[f"I{r}"], ws[f"K{r}"] = cat, total, 60, 60, 20
    t = wb.create_sheet("TARGET_AX_UNIL_NATI ")        # el nombre real tiene un espacio al final
    t.append(["OBJETIVOS"]); t.append([]); t.append([])
    t.append(["CAMPAÑA", "TIPO DE OBJETIVO", "COD.SUPERVISOR", "COD.EMPLEADO", "NOMBRE", "TARGET", "TARGET REMANENTE"])
    t.append(["DOVE 180 ML", "Cobertura", 3, 114, "X", 51, 0])
    t.append(["H`S SABORIZADA + H`S LIVIA", "Volumen", 3, 117, "Y", 1691, None])
    t.append(["REXONA AERO", "Volumen", 3, 119, "Z", 2740, None])
    t.append(["VIM BLOQUE DE MOCHILA", "Cobertura", 3, 102, "W", 10, 0])
    return wb


def test_lee_cobertura_y_cierra_con_el_total_de_la_distribuidora():
    df = OX.build_obj_cobertura(_libro_sintetico())
    assert set(df["categoria"]) == {"BPC", "FOOD", "HC"} and len(df) == 3 * 13
    for cat, total in (("BPC", 1398), ("FOOD", 1719), ("HC", 1680)):
        assert df[df.categoria == cat]["objetivo"].sum() == pytest.approx(total)
    bpc = df[df.categoria == "BPC"].set_index("vendedor_id")["objetivo"]
    assert bpc["101"] == pytest.approx((1398 - 60 - 60 - 20) / 9)      # 139,78 por preventista general
    assert bpc["100"] == bpc["119"] == 30 and bpc["120"] == 60 and bpc["122"] == 20


def test_lee_mis_ventas_en_formato_largo():
    df = OX.build_obj_mis_ventas(_libro_sintetico())
    assert sorted(df["campana"]) == ["DOVE_180ML", "HELLMANNS_SAB_LIV", "REXONA_AERO", "VIM_BLOQUE"]
    d = df.set_index("vendedor_id")
    assert d.loc["114", "tipo"] == "COBERTURA" and d.loc["114", "target"] == 51 and d.loc["114", "supervisor_id"] == "3"
    assert d.loc["119", "campana"] == "REXONA_AERO" and d.loc["119", "tipo"] == "VOLUMEN"


def test_campana_desconocida_falla():
    with pytest.raises(ValueError):
        OX.campana_clave("OTRA COSA")


def _hoja_campanas(filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campanas_candidatos"
    ws.append(["CAMPAÑA", "ARTICULO_ID", "DESCRIPCION", "LINEA", "ESTADO", "SUGERENCIA (S / N / ?)", "NOTA", "INCLUIR (S / N)"])
    for f in filas:
        ws.append(f)
    return wb


def test_cfg_campana_articulo_vacio_cuenta_como_si_y_lo_marcado_manda():
    wb = _hoja_campanas([
        ["DOVE 180 ML", "1996", "DOVE AC", "HAIR", "activo", "S", "", None],
        ["HELLMANNS SABORIZADA + LIVIANA", "1140", "MAY CLASICA", "DRESSING", "activo", "N", "", None],   # sugerencia N pero vacío
        ["REXONA AERO", "587", "CLINICAL", "DEO", "activo", "N", "", "n"],                                 # Juan marcó N
        ["VIM BLOQUE (pastilla / mochila)", "394", "VIM SAN", "HHC", "activo", "?", "", "S"],
        [None, None, None, None, None, None, None, None],
    ])
    df = OX.build_cfg_campana_articulo(wb).set_index("articulo_id")
    assert len(df) == 4
    assert df.loc["1140", "incluir"] and df.loc["1140", "fuente"] == "por defecto (S)"
    assert not df.loc["587", "incluir"] and df.loc["587", "fuente"] == "INCLUIR"
    assert df.loc["394", "campana"] == "VIM_BLOQUE" and df.loc["1996", "campana"] == "DOVE_180ML"


def test_cfg_campana_articulo_valor_invalido_falla():
    wb = _hoja_campanas([["DOVE 180 ML", "1", "X", "L", "activo", "S", "", "quizás"]])
    with pytest.raises(ValueError):
        OX.build_cfg_campana_articulo(wb)


def test_cfg_campana_articulo_formato_articulos_en_campana():
    """`articulos_en_campana.xlsx`: hoja única `Hoja1`, sin SUGERENCIA, y dos filas de campaña (VIM) con la misma clave."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(["CAMPAÑA", "ARTICULO_ID", "DESCRIPCION", "LINEA", "ESTADO", "INCLUIR (S / N)"])
    ws.append(["VIM BLOQUE (pastilla / mochila)", "1802", "VIM RIMBLOCK X6 FLORAL", "HHC", "activo", "N"])
    ws.append(["VIM BLOQUE (MOCHILA)", 2035, "VIM SAN PAST OCEANO (uxb:8)", "HHC", "activo", "S"])
    ws.append(["VIM BLOQUE (MOCHILA)", 393, "VIM SAN PAST DUO PACK OCEANO", "HHC", "activo", "S"])
    df = OX.build_cfg_campana_articulo(wb).set_index("articulo_id")
    assert set(df["campana"]) == {"VIM_BLOQUE"}
    assert list(df.index) == ["1802", "2035", "393"]            # ids numéricos del Excel pasan a texto
    assert df["incluir"].tolist() == [False, True, True]
    assert (df["sugerencia"] == "").all()


def test_build_obj_club_faro_cruza_contra_la_base_y_avisa_lo_que_no_esta():
    base = pd.DataFrame({"vendedor_id": ["120", "100", "201", "202"], "nombre": ["ACUÑA EMILIA", "GONZALBES MIGUEL", "NIEVA ALVARO", "PEREZ ROBERTO"]})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Distribuidora", "JEFE", "Supervisor", "DNI SUP", "Vendedores (SEGUN CUBO)", "DNI",
               "K+T - FAMILIA SMIRNOFF", "AS - BLANCOS DULCES (Cada sku suma 1 CC)", "AS - FAMILIA FRIZZE"])
    ws.append(["ZACARIAS PEJINAKIS PEÑAFLOR", "J", "Mauro", 1, "ACU?A EMILIA", 2, 29, 3, 1])           # exacta (? = Ñ)
    ws.append(["ZACARIAS PEJINAKIS PEÑAFLOR", "J", "Mauro", 1, "MIGUEL GONZALBEZ", 3, 1, 38, 31])      # aproximada (Z/S)
    ws.append(["ZACARIAS PEJINAKIS PEÑAFLOR", "J", "Natalia", 1, "ALVARO NIEVA", 4, 5, None, None])    # otro orden de nombre
    ws.append(["ZACARIAS PEJINAKIS PEÑAFLOR", "J", "Natalia", 1, "MAURO AGUIRRE", 6, 29, 3, 1])         # NO está en la base
    ws.append(["ZACARIAS PEJINAKIS PEÑAFLOR", "J", "Natalia", 1, "ROBERTO PERES", 7, 4, None, None])        # aproximada, NO confirmada
    ws.append(["ZACARIAS PEJINAKIS (JUJUY)", "J", "Fede", 1, "FLORES JOSE", 5, 32, 2, 5])
    df = OX.build_obj_club_faro(wb, dim_vendedor=base)
    assert "FLORES JOSE" not in set(df["nombre_excel"])                        # Jujuy no está en SIGMA
    por = df.drop_duplicates("nombre_excel").set_index("nombre_excel")
    assert por.loc["ACUÑA EMILIA", "vendedor_id"] == "120" and por.loc["ACUÑA EMILIA", "coincidencia"] == "exacta"
    assert por.loc["MIGUEL GONZALBEZ", "vendedor_id"] == "100" and por.loc["MIGUEL GONZALBEZ", "coincidencia"] == "confirmada"
    assert por.loc["MIGUEL GONZALBEZ", "nombre_base"] == "GONZALBES MIGUEL"
    assert por.loc["ALVARO NIEVA", "vendedor_id"] == "201" and por.loc["ALVARO NIEVA", "coincidencia"] == "exacta"
    assert por.loc["MAURO AGUIRRE", "vendedor_id"] == "" and por.loc["MAURO AGUIRRE", "coincidencia"] == "sin coincidencia"
    emilia = df[df["vendedor_id"] == "120"].set_index("linea")
    assert emilia.loc["SMIRNOFF", "objetivo"] == 29 and emilia.loc["SMIRNOFF", "tipo_cliente"] == "TRAD"
    assert emilia.loc["FRIZZE", "tipo_cliente"] == "AS" and not emilia["supuesto"].any()
    assert not df[df["vendedor_id"] == "100"]["supuesto"].any()                # Gonzalbes: cruce confirmado por Juan
    assert por.loc["ROBERTO PERES", "vendedor_id"] == "202" and por.loc["ROBERTO PERES", "coincidencia"] == "aproximada"
    assert df[df["vendedor_id"] == "202"]["supuesto"].all()                    # aproximada sin confirmar = marcada para revisar
    assert len(df[df["nombre_excel"] == "ALVARO NIEVA"]) == 1                  # objetivos vacíos no generan fila


def test_build_obj_club_faro_exige_la_base():
    wb = openpyxl.Workbook()
    wb.active.append(["Distribuidora", "Supervisor", "Vendedores", "K+T - FAMILIA SMIRNOFF"])
    with pytest.raises(ValueError):
        OX.build_obj_club_faro(wb)


def test_build_cfg_club_faro_articulo():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(["LINEA CLUB FARO", "ARTICULO_ID", "DESCRIPCION", "LINEA SIGMA", "GRUPO", "ESTADO", "INCLUIR (S / N)", "NOTA"])
    ws.append(["K+T - FAMILIA SMIRNOFF", "1162", "SMIRNOFF TROPICAL", "SPIRITS", "", "activo", None, ""])
    ws.append(["K+T - FAMILIA SMIRNOFF", "8005", "RECONOCIMIENTO", "CON ALCOHOL", "", "activo", "N", ""])
    ws.append(["AS - BLANCOS DULCES (cada SKU suma 1)", 1930, "ALMA MORA", "VINO", "", "activo", "s", ""])
    df = OX.build_cfg_club_faro_articulo(wb).set_index("articulo_id")
    assert df.loc["1162", "incluir"] and df.loc["1162", "fuente"] == "por defecto (S)"
    assert not df.loc["8005", "incluir"] and df.loc["1930", "linea"] == "BLANCOS_DULCES" and df.loc["1930", "fuente"] == "INCLUIR"


def test_vendedores_fuera_de_la_base():
    base = pd.DataFrame({"vendedor_id": ["100", "101", "102"], "nombre": ["A", "B", "C"], "activo": [True, True, False]})
    tablas = {"obj_mis_ventas": pd.DataFrame({"vendedor_id": ["100", "999", "100"]}),
              "cfg_vendedor_perfil": pd.DataFrame({"vendedor_id": ["101", "102"]}),
              "obj_facturacion": pd.DataFrame({"escalon": [1]})}                    # sin vendedor_id: se ignora
    df = OX.vendedores_fuera_de_la_base(tablas, base)
    assert set(zip(df["tabla"], df["vendedor_id"], df["motivo"])) == {
        ("obj_mis_ventas", "999", "no está en la base"), ("cfg_vendedor_perfil", "102", "desactivado en la base")}
    assert OX.vendedores_fuera_de_la_base({"x": pd.DataFrame({"vendedor_id": ["100"]})}, base).empty


def test_build_cfg_11_titulares_articulo():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Articulos"
    ws.append(["LINEA 11 TITULARES", "ARTICULO_ID", "DESCRIPCION", "INCLUIR (S / N)"])
    ws.append(["Alma Mora", "1439", "ALMA MORA RVA", None])
    ws.append(["Trapiche Reserva", 1500, "TRAPICHE RESERVA", "s"])
    ws.append(["Dada", "2000", "DADA 1", "N"])
    df = OX.build_cfg_11_titulares_articulo(wb).set_index("articulo_id")
    assert df.loc["1439", "linea"] == "ALMA_MORA" and bool(df.loc["1439", "incluir"]) and df.loc["1439", "fuente"] == "por defecto (S)"
    assert bool(df.loc["1500", "incluir"]) and df.loc["1500", "fuente"] == "INCLUIR"
    assert not bool(df.loc["2000", "incluir"])
    ws.append(["Alaris", "9", "X", "quizás"])
    with pytest.raises(ValueError):
        OX.build_cfg_11_titulares_articulo(wb)
