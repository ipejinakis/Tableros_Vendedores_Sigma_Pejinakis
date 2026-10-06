"""Genera `articulos_11_titulares.xlsx` con los CANDIDATOS de 11 Titulares para que Juan marque a mano.

Hoja `Articulos`: artículos de la base (dim_articulo, incluye desactivados) que parecen de cada una de las 11 líneas
titulares (por palabras en la descripción), con INCLUIR vacío para marcar S / N. Hoja `Rubros`: rubros de clientes de la
base con la cantidad de clientes y el CANAL sugerido (Autoservicios / Tradicionales / OP y VTK), a confirmar.
No toma nada de los Excel: todo sale de la base. No escribe en la base.

Uso (desde la raíz del repo, venv activo):
    python scripts/etl/generar_11_titulares.py
    python scripts/etl/generar_11_titulares.py --salida "../articulos_11_titulares.xlsx"
"""
import argparse
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from sigma_conn import negocio as N  # noqa: E402
from sigma_conn import tableros as TB  # noqa: E402

# (clave, nombre de la línea como en el reporte de Peñaflor, patrón de búsqueda sobre la descripción, patrón a excluir)
LINEAS = [
    ("ALMA_MORA", "Alma Mora", r"ALMA MORA", None),
    ("TRAPICHE_RESERVA", "Trapiche Reserva", r"TRAPICHE", None),
    ("FINCA_LAS_MORAS", "Finca Las Moras", r"LAS MORAS|F\.LAS MORAS", None),
    ("ALARIS", "Alaris", r"ALARIS", None),
    ("DON_DAVID", "Don David", r"DON DAVID", None),
    ("DADA", "Dada", r"\bDADA\b", None),
    ("SMIRNOFF_FLAVORS", "Smirnoff Flavors", r"SMIRNOFF", r"\bICE\b|LATA|RECONOCIMIENTO|\bBC\b"),
    ("LOS_ARBOLES", "Los Arboles", r"LOS ARBOLES", None),
    ("ANTARES", "Antares", r"ANTARES", None),
    ("SMIRNOFF_ICE", "Smirnoff Ice", r"SMIRNOFF ICE", None),
    ("GORDONS_FLAVOURS", "Gordons Flavours", r"GORDON", None),
]
CANALES = ("Autoservicios", "Tradicionales", "OP y VTK (On Premise, Vinotecas, Tienda de Bebidas, Catering)")

ap = argparse.ArgumentParser()
ap.add_argument("--salida", default=str(ROOT.parent / "articulos_11_titulares.xlsx"))
a = ap.parse_args()

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")


def sin_tildes(s) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(s).upper()) if unicodedata.category(c) != "Mn")


def limpio(v):
    """Valor apto para Excel: los vacíos de pandas (<NA>, NaN) pasan a None; los tipos de numpy a Python."""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    return v.item() if hasattr(v, "item") else v


store = TB.abrir_store()
art = store.read_table("dim_articulo")
cli = store.read_table("dim_cliente")
art = art.drop_duplicates("articulo_id").assign(_d=lambda d: d["descripcion"].map(sin_tildes))

filas = []
for clave, nombre, pat, excl in LINEAS:
    m = art["_d"].str.contains(pat, regex=True, na=False)
    if excl:
        m &= ~art["_d"].str.contains(excl, regex=True, na=False)
    for r in art[m].sort_values("articulo_id").itertuples():
        desact = str(getattr(r, "desactivado", "")).lower() == "true"
        filas.append((nombre, str(r.articulo_id), r.descripcion, getattr(r, "proveedor", ""), getattr(r, "linea", ""),
                      getattr(r, "grupo", ""), "desactivado" if desact else "activo",
                      getattr(r, "unidades_por_bulto", None), None, ""))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Articulos"
ws.append(["LINEA 11 TITULARES", "ARTICULO_ID", "DESCRIPCION", "PROVEEDOR", "LINEA SIGMA", "GRUPO SIGMA", "ESTADO",
           "UNIDADES POR BULTO", "INCLUIR (S / N)", "NOTA"])
for f in filas:
    ws.append([limpio(x) for x in f])
amarillo = PatternFill("solid", fgColor="FFF2CC")
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
    for c in row:
        c.fill = amarillo
for col, w in zip("ABCDEFGHIJ", (24, 12, 42, 22, 16, 22, 12, 12, 14, 40)):
    ws.column_dimensions[col].width = w

wr = wb.create_sheet("Rubros")
wr.append(["RUBRO_COD", "RUBRO", "CLIENTES", "CANAL (confirmar)"])
rub = cli.assign(rubro_cod=cli["rubro_cod"].astype(str)).groupby(["rubro_cod", "rubro"], dropna=False).size().reset_index(name="n")
for r in rub.sort_values("n", ascending=False).itertuples():
    nombre = sin_tildes(r.rubro)
    if r.rubro_cod in N.RUBROS_AS:
        sug = CANALES[0]
    elif "VINOTECA" in nombre or "TIENDA DE BEBIDAS" in nombre or "ON PREMISE" in nombre or "CATERING" in nombre:
        sug = CANALES[2]
    else:
        sug = CANALES[1]
    wr.append([limpio(r.rubro_cod), limpio(r.rubro), int(r.n), sug])
for c in wr[1]:
    c.font = Font(bold=True)
for row in wr.iter_rows(min_row=2, min_col=4, max_col=4):
    for c in row:
        c.fill = amarillo
for col, w in zip("ABCD", (10, 30, 10, 60)):
    wr.column_dimensions[col].width = w

# Atributos personalizados de los clientes en SIGMA (columnas attr_* de dim_cliente; requieren el ETL completo).
wa = wb.create_sheet("Atributos")
wa.append(["ATRIBUTO", "VALOR", "CLIENTES", "CON RUBRO MAS FRECUENTE"])
atributos = [c for c in cli.columns if c.startswith("attr_")]
if not atributos:
    print("AVISO: dim_cliente todavía no tiene los atributos de SIGMA (attr_*). Correr el ETL completo: python scripts/etl/run_etl.py")
for c in atributos:
    vc = cli[c].dropna().value_counts().head(12)
    for valor, n in vc.items():
        rub_top = cli.loc[cli[c] == valor, "rubro"].astype(str).value_counts().head(3)
        wa.append([c.removeprefix("attr_"), limpio(valor), int(n), "; ".join(f"{k} ({v})" for k, v in rub_top.items())])
for c in wa[1]:
    c.font = Font(bold=True)
for col, w in zip("ABCD", (26, 30, 10, 80)):
    wa.column_dimensions[col].width = w

wl = wb.create_sheet("Leeme")
for t in ("Candidatos de 11 Titulares (Peñaflor), generados desde la base de datos. Nada de esto está confirmado.",
          "Hoja Articulos: marcá INCLUIR con S (cuenta para esa línea) o N (no cuenta). Un artículo puede estar en una sola línea.",
          "Si falta un artículo, sumá una fila al final con el nombre exacto de la línea y el ARTICULO_ID de SIGMA.",
          "UNIDADES POR BULTO sirve para la regla: autoservicio suma con 1 caja/bulto; tradicional con 3 unidades iguales.",
          "Hoja Atributos: atributos propios de los clientes en SIGMA (p. ej. `vinotecas` = \"11 titulares vinotec\"): sirven para ver si alguno define el canal de 11 Titulares.",
          "Hoja Rubros: el CANAL es una sugerencia (autoservicio = rubros 01, 98, 70; vinotecas/tienda de bebidas/on premise/catering = OP y VTK; el resto = Tradicionales). Corregilo donde haga falta.",
          "Después de editar, avisame (no hace falta correr nada)."):
    wl.append([t])
wl.column_dimensions["A"].width = 130
wb.save(a.salida)

print(f"Archivo: {a.salida}")
print(f"Artículos candidatos: {len(filas)}")
for _, nombre, _, _ in LINEAS:
    n = sum(1 for f in filas if f[0] == nombre)
    print(f"  {nombre}: {n}" + ("   <-- SIN CANDIDATOS: revisar" if n == 0 else ""))
print(f"Rubros de clientes: {len(rub)}")
